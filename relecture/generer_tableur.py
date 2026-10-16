"""
Génère le tableur de relecture médicale à partir de data/syndromik_data.json.

    python3 relecture/generer_tableur.py

Produit relecture/syndromik_relecture_v<version>.xlsx :
  - « Mode d'emploi »  : consignes + exemple de ligne remplie
  - « Synthèse »       : avancement de la relecture (formules)
  - « À valider »      : ajouts proposés ("a_valider": true), à traiter en premier
  - une feuille par type (Signes, Syndromes, Pathologies, Spécialités) avec leurs liens
"""
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
data = json.loads((ROOT / "data" / "syndromik_data.json").read_text(encoding="utf-8"))
version = data["meta"]["version"]
OUT = ROOT / "relecture" / f"syndromik_relecture_v{version}.xlsx"

DECISIONS = ["Oui", "Non", "À modifier"]
URGENCES = {"vitale": "Vitale", "elevee": "Élevée", "moderee": "Modérée", "faible": "Faible"}

FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="0E5C63")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_MUTED = Font(name=FONT, size=9, italic=True, color="5C6B6A")
FILL_HEAD = PatternFill("solid", fgColor="0E5C63")
FILL_INPUT = PatternFill("solid", fgColor="FFF6D6")   # cellules à remplir par le relecteur
FILL_AV = PatternFill("solid", fgColor="FCEFD2")      # ligne d'un ajout « à valider »
THIN = Side(style="thin", color="E1E8E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# --- Index -----------------------------------------------------------------
TYPES = [("signes", "Signe"), ("syndromes", "Syndrome"),
         ("pathologies", "Pathologie"), ("specialites", "Spécialité")]
nom = {e["id"]: e["nom"] for t, _ in TYPES for e in data[t]}
voisins = defaultdict(lambda: defaultdict(list))  # id -> type voisin -> [(nom, a_valider)]
L = data["liens"]
for cle, a, b in [("signe_syndrome", "signe", "syndrome"),
                  ("syndrome_pathologie", "syndrome", "pathologie"),
                  ("pathologie_specialite", "pathologie", "specialite")]:
    for x in L.get(cle, []):
        av = bool(x.get("a_valider"))
        voisins[x[a]][b].append((nom[x[b]], av))
        voisins[x[b]][a].append((nom[x[a]], av))


def liste(entite_id, type_voisin):
    items = sorted(voisins[entite_id][type_voisin], key=lambda v: v[0])
    return "\n".join(n + (" [à valider]" if av else "") for n, av in items)


# --- Mise en forme ---------------------------------------------------------
def entete(ws, row, cols):
    for i, (titre, largeur) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=titre)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, WRAP, BORDER
        ws.column_dimensions[c.column_letter].width = largeur
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{ws.cell(row=row, column=len(cols)).column_letter}{row}"


def ligne(ws, row, valeurs, n_input, fill=None):
    """Écrit une ligne ; les n_input dernières colonnes sont à remplir."""
    total = len(valeurs) + n_input
    for i in range(1, total + 1):
        c = ws.cell(row=row, column=i, value=valeurs[i - 1] if i <= len(valeurs) else None)
        c.font, c.alignment, c.border = F_BODY, WRAP, BORDER
        if i > len(valeurs):
            c.fill = FILL_INPUT
        elif fill:
            c.fill = fill


def validation_decision(ws, col, first, last):
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Décision", error="Choisir Oui, Non ou À modifier.")
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")


wb = Workbook()

# --- Mode d'emploi ---------------------------------------------------------
ws = wb.active
ws.title = "Mode d'emploi"
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 90
ws["A1"] = f"Syndromik — relecture médicale de la base v{version}"
ws["A1"].font = F_TITLE
lignes = [
    ("Objectif", "Vérifier que le contenu est conforme aux protocoles nationaux ivoiriens "
                 "avant tout usage réel de l'application."),
    ("Par où commencer", "Feuille « À valider » : ajouts proposés non relus "
                         f"({sum(1 for s in data['signes'] if s.get('a_valider'))} signes et "
                         f"{sum(1 for x in L['signe_syndrome'] if x.get('a_valider'))} liens). "
                         "Ensuite, si possible, les feuilles par type (contenu existant)."),
    ("Cellules à remplir", "Uniquement les colonnes sur fond jaune pâle : « Décision » et « Commentaire »."),
    ("Décision", "Oui = correct tel quel · Non = à supprimer · À modifier = préciser la "
                 "correction dans « Commentaire »."),
    ("Urgence", "Feuille « Pathologies » : vérifier le niveau d'urgence ; en cas de désaccord, "
                "choisir « À modifier » et indiquer le bon niveau (Vitale, Élevée, Modérée, Faible)."),
    ("[à valider]", "Dans les listes de liens, ce repère signale un lien ajouté et non encore relu."),
    ("Avancement", "Feuille « Synthèse » : se met à jour automatiquement."),
    ("Ne pas modifier", "Les colonnes ID, Nom, Description, etc. (sauf demande explicite). "
                        "Les corrections passent par « Commentaire »."),
]
r = 3
for k, v in lignes:
    ws.cell(row=r, column=1, value=k).font = F_BOLD
    c = ws.cell(row=r, column=2, value=v)
    c.font, c.alignment = F_BODY, WRAP
    ws.cell(row=r, column=1).alignment = WRAP
    r += 1
r += 1
ws.cell(row=r, column=1, value="Exemple de ligne remplie").font = F_BOLD
r += 1
ex_cols = ["Type", "Élément", "Détail", "Décision", "Commentaire"]
ex_val = ["Lien", "Fièvre → Syndrome ictérique",
          "Le signe « Fièvre » est ajouté au syndrome « Syndrome ictérique ».",
          "À modifier",
          "Fièvre présente dans le palu, la leptospirose et la fièvre jaune, mais pas dans "
          "l'ictère néonatal : OK pour le lien, à nuancer dans la description."]
for i, (h, v) in enumerate(zip(ex_cols, ex_val)):
    ws.cell(row=r + i, column=1, value=h).font = F_BOLD
    c = ws.cell(row=r + i, column=2, value=v)
    c.font, c.alignment = F_BODY, WRAP
    if h in ("Décision", "Commentaire"):
        c.fill = FILL_INPUT
r += len(ex_cols) + 1
ws.cell(row=r, column=1, value="Source").font = F_BOLD
c = ws.cell(row=r, column=2, value=f"Généré automatiquement depuis data/syndromik_data.json "
                                    f"(version {version}) par relecture/generer_tableur.py.")
c.font, c.alignment = F_MUTED, WRAP

# --- À valider -------------------------------------------------------------
ws = wb.create_sheet("À valider")
cols = [("Type", 10), ("ID", 26), ("Élément", 32), ("Détail", 60),
        ("Décision", 13), ("Commentaire", 50)]
entete(ws, 1, cols)
r = 2
for s in data["signes"]:
    if s.get("a_valider"):
        detail = s["description"]
        if s.get("synonymes"):
            detail += "\nSynonymes : " + ", ".join(s["synonymes"])
        ligne(ws, r, ["Nouveau signe", s["id"], s["nom"], detail], 2)
        r += 1
for cle, a, b in [("signe_syndrome", "signe", "syndrome"),
                  ("syndrome_pathologie", "syndrome", "pathologie"),
                  ("pathologie_specialite", "pathologie", "specialite")]:
    for x in L.get(cle, []):
        if x.get("a_valider"):
            ligne(ws, r, ["Lien", f"{x[a]} → {x[b]}", f"{nom[x[a]]} → {nom[x[b]]}",
                          f"Le {a} « {nom[x[a]]} » est ajouté au {b} « {nom[x[b]]} »."], 2)
            r += 1
last_av = max(r - 1, 2)
validation_decision(ws, "E", 2, last_av)

# --- Feuilles par type -----------------------------------------------------
feuilles = {}
specs = {
    "signes": [("Syndromes liés", "syndrome")],
    "syndromes": [("Signes", "signe"), ("Pathologies", "pathologie")],
    "pathologies": [("Syndromes", "syndrome"), ("Spécialités", "specialite")],
    "specialites": [("Pathologies", "pathologie")],
}
for t, label in TYPES:
    ws = wb.create_sheet(label + "s" if label != "Spécialité" else "Spécialités")
    cols = [("ID", 24), ("Nom", 28), ("Synonymes", 26), ("Description", 55)]
    if t == "pathologies":
        cols.append(("Urgence", 11))
    cols += [(titre, 34) for titre, _ in specs[t]]
    cols += [("Décision", 13), ("Commentaire", 45)]
    entete(ws, 1, cols)
    r = 2
    for e in sorted(data[t], key=lambda e: e["nom"]):
        vals = [e["id"], e["nom"] + (" [à valider]" if e.get("a_valider") else ""),
                ", ".join(e.get("synonymes", [])), e.get("description", "")]
        if t == "pathologies":
            vals.append(URGENCES.get(e.get("urgence"), e.get("urgence") or ""))
        vals += [liste(e["id"], tv) for _, tv in specs[t]]
        ligne(ws, r, vals, 2, FILL_AV if e.get("a_valider") else None)
        r += 1
    dec_col = ws.cell(row=1, column=len(cols) - 1).column_letter
    validation_decision(ws, dec_col, 2, r - 1)
    feuilles[ws.title] = (dec_col, 2, r - 1)
feuilles = {"À valider": ("E", 2, last_av), **feuilles}

# --- Synthèse (formules) ---------------------------------------------------
ws = wb.create_sheet("Synthèse", 1)
ws["A1"] = "Avancement de la relecture"
ws["A1"].font = F_TITLE
cols = [("Feuille", 18), ("Éléments", 11), ("Oui", 9), ("Non", 9),
        ("À modifier", 12), ("Restants", 11), ("Avancement", 13)]
for i, (h, w) in enumerate(cols, start=1):
    c = ws.cell(row=3, column=i, value=h)
    c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
    ws.column_dimensions[c.column_letter].width = w
r = 4
for nom_f, (col, a, b) in feuilles.items():
    rng = f"'{nom_f}'!${col}${a}:${col}${b}"
    vals = [nom_f, f"=ROWS({rng})",
            f'=COUNTIF({rng},"Oui")', f'=COUNTIF({rng},"Non")', f'=COUNTIF({rng},"À modifier")',
            f"=B{r}-C{r}-D{r}-E{r}", f"=IF(B{r}=0,0,(C{r}+D{r}+E{r})/B{r})"]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font, c.border = F_BODY, BORDER
    ws.cell(row=r, column=7).number_format = "0%"
    r += 1
tot = ["Total"] + [f"=SUM({ws.cell(row=4, column=i).column_letter}4:"
                   f"{ws.cell(row=r - 1, column=i).column_letter}{r - 1})" for i in range(2, 7)]
tot.append(f"=IF(B{r}=0,0,(C{r}+D{r}+E{r})/B{r})")
for i, v in enumerate(tot, start=1):
    c = ws.cell(row=r, column=i, value=v)
    c.font, c.border = F_BOLD, BORDER
ws.cell(row=r, column=7).number_format = "0%"
ws.cell(row=r + 2, column=1,
        value="Calculé à partir des colonnes « Décision » de chaque feuille.").font = F_MUTED

for w in wb.worksheets:
    w.sheet_view.showGridLines = w.title not in ("Mode d'emploi", "Synthèse")
wb.save(OUT)
print(OUT)
